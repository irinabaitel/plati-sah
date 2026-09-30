"""Exportă sumele din foaia „Evidență plăți" pentru pagina plati.html.

Rulare:  python plati_export.py
Citește cel mai recent __SAH 2627__pilot*.xlsx din OneDrive (pilot, pilot_v1, ...),
dintr-o copie, ca să meargă și cu Excel deschis,
și scrie plati_date.js. Numele copiilor NU apar în clar în fișier: fiecare copil e găsit
după amprenta SHA-256 a numelui normalizat (fără diacritice, cuvintele în ordine alfabetică).

Plățile: se citesc toate exporturile CSV din ING Business puse în __SAH_26-27__\\incasari\\
(„Rapoarte → Extras de cont CSV”; se pot suprapune, dublurile se elimină după „Referinta bancii”).
Fiecare încasare se potrivește cu copilul după numele din detalii + numele plătitorului;
suma pentru frați se împarte după cât are fiecare de plată. Când există CSV-uri, „Achitat” pe
site vine DOAR din bancă; în Excel coloanele „Achitat” sunt o copie: scriptul scrie
achitat_de_lipit_<luna>.txt (o valoare pe rând, în ordinea rândurilor din Evidență plăți),
iar userul o lipește la rândul 7 al coloanei lunii. Nu scriem direct în .xlsx (openpyxl strică
formatările condiționate și validările).
Datele bancare rămân în OneDrive: potriviri.json (corecturi + conturi învățate) și
raport_potriviri.txt (ce plată la ce copil a mers). În plati_date.js ajung doar sumele.
"""
import csv, glob, hashlib, itertools, json, re, shutil, tempfile, unicodedata, datetime, os

FOLDER = os.path.expanduser(r'~\OneDrive\__SAH_26-27__')
TIPAR = '__SAH 2627__pilot*.xlsx'   # fișierele din „arhiva versiuni vechi” nu intră (subfolder)
INCASARI = os.path.join(FOLDER, 'incasari')
POTRIVIRI = os.path.join(INCASARI, 'potriviri.json')
RAPORT = os.path.join(INCASARI, 'raport_potriviri.txt')
# luna calendaristică -> coloana din tabel (sep = 0 ... iun = 9)
LUNA_IDX = {9: 0, 10: 1, 11: 2, 12: 3, 1: 4, 2: 5, 3: 6, 4: 7, 5: 8, 6: 9}
DEST = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'plati_date.js')
SARE = 'lab-sah-2627'   # trebuie să fie identică cu cea din plati.html

LUNI = ['septembrie', 'octombrie', 'noiembrie', 'decembrie', 'ianuarie',
        'februarie', 'martie', 'aprilie', 'mai', 'iunie']
PRIMA_COL = 4           # D = Prez. septembrie; fiecare lună are 3 coloane (Prez., De plată, Achitat)
PRIMUL_RAND = 7


def normalizeaza(text):
    text = unicodedata.normalize('NFD', str(text)).encode('ascii', 'ignore').decode()
    text = ''.join(c if c.isalpha() else ' ' for c in text.upper())
    return text.split()


def chei(nume):
    """Amprente pentru numele complet și pentru orice combinație de cel puțin 2 cuvinte
    (ca „Iacovici Iris" să-l găsească și pe „IACOVICI IRIS MIHAELA")."""
    cuvinte = normalizeaza(nume)
    rez = set()
    for k in range(min(2, len(cuvinte)), len(cuvinte) + 1):
        for comb in itertools.combinations(cuvinte, k):
            rez.add(hashlib.sha256((SARE + ' '.join(sorted(comb))).encode()).hexdigest()[:20])
    return rez


def numar(v):
    return v if isinstance(v, (int, float)) else 0


def la_fel(a, b):
    """Cuvinte egale sau cu o singură literă diferită (Patrick / Patrik)."""
    if a == b:
        return True
    if min(len(a), len(b)) < 4 or abs(len(a) - len(b)) > 1:
        return False
    prev = list(range(len(b) + 1))
    for i, ca in enumerate(a, 1):
        cur = [i]
        for j, cb in enumerate(b, 1):
            cur.append(min(prev[j] + 1, cur[j - 1] + 1, prev[j - 1] + (ca != cb)))
        prev = cur
    return prev[-1] <= 1


def citeste_incasari():
    """Toate încasările din CSV-urile ING, fără dubluri."""
    vazute, rez = set(), []
    for f in sorted(glob.glob(os.path.join(INCASARI, '*.csv'))):
        with open(f, encoding='utf-8-sig', errors='replace') as fh:
            for r in csv.reader(fh, delimiter=';'):
                if len(r) < 10 or r[0].startswith('numar cont'):
                    continue
                suma = float(r[2].replace('.', '').replace(',', '.'))
                if suma <= 0 or r[4].strip() != 'Incasare':
                    continue
                m = re.search(r'Referinta bancii\s*(\S+)', r[9])
                ref = m.group(1) if m else f'{r[1]}|{r[2]}|{r[7]}|{r[9]}'
                if ref in vazute:
                    continue
                vazute.add(ref)
                rez.append({'ref': ref, 'data': datetime.datetime.strptime(r[1], '%d.%m.%Y').date(),
                            'suma': suma, 'platitor': r[5].strip(), 'iban': r[7].strip(),
                            'detalii': re.sub(r'\s*Referinta bancii.*', '', r[9]).strip()})
    return sorted(rez, key=lambda p: p['data'])


def potriveste(plata, nume_copii, potriviri):
    """Întoarce (lista de nume, cum s-a găsit)."""
    manual = potriviri['referinta'].get(plata['ref'])
    if manual is not None:
        return manual, 'corectat de mână'
    text = normalizeaza(plata['detalii'] + ' ' + plata['platitor'])
    gasiti = []
    for nume in nume_copii:
        cuv = normalizeaza(nume)
        potrivite = sum(any(la_fel(c, t) for t in text) for c in cuv)
        if potrivite == len(cuv) or potrivite >= 2:
            gasiti.append(nume)
    if len({normalizeaza(n)[0] for n in gasiti}) > 1:
        # mai mulți copii se împart o plată doar dacă sunt frați; altfel e un nume comun
        # („Cojocaru Vlad Mihai” nu e și pentru MIHAI VLAD) -> rămâne cine are numele plătitorului
        platitor = normalizeaza(plata['platitor'])
        gasiti = [n for n in gasiti if any(la_fel(normalizeaza(n)[0], t) for t in platitor)]
        if len({normalizeaza(n)[0] for n in gasiti}) != 1:
            return [], 'NEGĂSIT (se potrivește cu mai mulți copii)'
    if gasiti:
        return gasiti, 'după nume'
    memorat = potriviri['iban'].get(plata['iban'])
    if memorat:
        return memorat, 'după contul plătitorului'
    return [], 'NEGĂSIT'


def aplica_incasari(nume_copii, copii, curenta):
    if not os.path.isdir(INCASARI) or not glob.glob(os.path.join(INCASARI, '*.csv')):
        return False
    for luni in copii:           # plățile vin din bancă; ce e în Excel e doar copia lor
        for l in luni:
            l[2] = 0
    try:
        with open(POTRIVIRI, encoding='utf-8') as f:
            potriviri = json.load(f)
    except FileNotFoundError:
        potriviri = {}
    potriviri.setdefault('referinta', {})
    potriviri.setdefault('iban', {})
    poz = {n: i for i, n in enumerate(nume_copii)}

    linii, total, negasite = [], 0, 0
    for p in citeste_incasari():
        cine, cum = potriveste(p, nume_copii, potriviri)
        cine = [n for n in cine if n in poz]
        antet = f"{p['data']:%d.%m}  {p['suma']:>7.2f}  {p['platitor']} — „{p['detalii']}”"
        if not cine:
            negasite += 1
            linii.append(f'{antet}\n      ?? NEGĂSIT — spune-i lui Claude al cui copil e')
            continue
        if cum == 'după nume' and p['iban']:
            potriviri['iban'][p['iban']] = cine      # luna viitoare îl recunoaște și după cont
        luna = LUNA_IDX.get(p['data'].month, 0 if p['data'].month in (7, 8) else 9)
        # frații: suma se împarte după cât are fiecare de plată până acum
        datorii = [sum(l[1] for l in copii[poz[n]][:curenta + 1]) for n in cine]
        baza = sum(datorii) or len(cine)
        parti = [round(p['suma'] * (d if sum(datorii) else 1) / baza) for d in datorii]
        parti[0] += round(p['suma'] - sum(parti), 2)
        for n, s in zip(cine, parti):
            copii[poz[n]][luna][2] += s
        total += p['suma']
        linii.append(f'{antet}\n      -> ' + ', '.join(f'{n} ({s:g} lei)' for n, s in zip(cine, parti)) + f'  [{cum}]')

    with open(POTRIVIRI, 'w', encoding='utf-8') as f:
        json.dump(potriviri, f, ensure_ascii=False, indent=1)
    with open(RAPORT, 'w', encoding='utf-8') as f:
        f.write(f'Potrivirea încasărilor — {datetime.datetime.now():%d.%m.%Y %H:%M}\n\n' + '\n'.join(linii) + '\n')
    print('\n'.join(linii))
    print(f'Încasări atribuite: {total:g} lei; negăsite: {negasite}')
    return True


def scrie_de_lipit(randuri, copii, prim, ultim):
    """Coloana „Achitat” pentru fiecare lună cu plăți, rând cu rând de la PRIMUL_RAND."""
    for f in glob.glob(os.path.join(INCASARI, 'achitat_de_lipit_*.txt')):
        os.remove(f)
    pe_rand = dict(zip(randuri, copii))
    for i, luna in enumerate(LUNI):
        if not any(c[i][2] for c in copii):
            continue
        col = openpyxl_col(PRIMA_COL + 3 * i + 2)
        # 0 (nu gol) la cine n-a plătit: Excel ignoră rândurile goale de la final la lipire,
        # iar așa se suprascrie mereu toată coloana, până la ultimul copil
        valori = [(f'{pe_rand[r][i][2]:g}' if r in pe_rand else '')
                  for r in range(prim, ultim + 1)]
        cale = os.path.join(INCASARI, f'achitat_de_lipit_{luna}.txt')
        with open(cale, 'w', encoding='utf-8', newline='') as f:
            f.write('\r\n'.join(valori))
        print(f'De lipit: {luna} -> coloana {col}, începând cu {col}{prim} ({cale})')


def openpyxl_col(n):
    from openpyxl.utils import get_column_letter
    return get_column_letter(n)


def main():
    import openpyxl
    fisiere = [f for f in glob.glob(os.path.join(FOLDER, TIPAR))
               if not os.path.basename(f).startswith('~$')]   # ~$... = fișierul de blocare al Excel
    SURSA = max(fisiere, key=os.path.getmtime)
    print('Sursa:', os.path.basename(SURSA))
    copie = os.path.join(tempfile.gettempdir(), 'plati_export_copie.xlsx')
    try:
        shutil.copyfile(SURSA, copie)
    except PermissionError:   # Excel ține fișierul blocat; Copy-Item din PowerShell reușește
        import subprocess
        subprocess.run(['powershell', '-NoProfile', '-Command',
                        f'Copy-Item -LiteralPath "{SURSA}" -Destination "{copie}" -Force'], check=True)
    ws = openpyxl.load_workbook(copie, data_only=True)['Evidență plăți']
    tarif = ws['C3'].value

    copii, index, nume_copii, randuri = [], {}, [], []
    ultimul = ws.max_row
    for r in range(PRIMUL_RAND, ws.max_row + 1):
        nume = ws.cell(r, 2).value
        if str(nume).strip().upper() == 'TOTAL':
            ultimul = r - 1
            break
        if not nume:
            continue
        randuri.append(r)
        luni = []
        for i in range(len(LUNI)):
            c = PRIMA_COL + 3 * i
            luni.append([numar(ws.cell(r, c).value), numar(ws.cell(r, c + 1).value),
                         numar(ws.cell(r, c + 2).value)])
        idx = len(copii)
        copii.append(luni)
        nume_copii.append(str(nume).strip())
        for k in chei(nume):
            index.setdefault(k, []).append(idx)

    # luna curentă = ultima lună în care a venit cineva
    curenta = max((i for i in range(len(LUNI)) if any(c[i][0] for c in copii)), default=0)
    # copiii scutiți de plată (lista e în potriviri.json, privat): pe site apar cu 0 de plată
    try:
        with open(POTRIVIRI, encoding='utf-8') as f:
            scutiti = {' '.join(normalizeaza(n)) for n in json.load(f).get('scutiti', [])}
    except FileNotFoundError:
        scutiti = set()
    for n, luni in zip(nume_copii, copii):
        if ' '.join(normalizeaza(n)) in scutiti:
            for l in luni:
                l[0] = l[1] = 0
            print('Scutit de plată:', n)

    if aplica_incasari(nume_copii, copii, curenta):
        scrie_de_lipit(randuri, copii, PRIMUL_RAND, ultimul)

    date = {
        'actualizat': datetime.date.today().strftime('%d.%m.%Y'),
        'tarif': tarif,
        'luni': LUNI,
        'lunaCurenta': curenta,
        'copii': copii,
        'index': index,
    }
    with open(DEST, 'w', encoding='utf-8') as f:
        f.write('// Generat de plati_export.py — nu modifica de mână.\n')
        f.write('window.PLATI = ' + json.dumps(date, ensure_ascii=False, separators=(',', ':')) + ';\n')
    print(f'{len(copii)} copii exportați, luna curentă: {LUNI[curenta]}, tarif {tarif} lei -> {DEST}')


if __name__ == '__main__':
    main()
